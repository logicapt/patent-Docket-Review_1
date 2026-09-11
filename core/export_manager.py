import io
import pandas as pd
from typing import List, Dict, Any
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def evidence_columns(item):
    """The CSV and workbook carry the same source and date provenance."""
    return {
        "Status": item.get("status_display", ""),
        "Litigation Stage": item.get("litigation_stage", ""),
        "Invalidity Contentions Observed?": "Yes" if item.get("has_invalidity_contentions") else "Not observed",
        "Verification Status": item.get("verification_status", "UNVERIFIED"),
        "Verification Source": item.get("source_used", "No verified source"),
        "Verification Details": item.get("verification_message", ""),
        "Import File": item.get("import_filename", ""),
        "Imported At (UTC)": item.get("imported_at", ""),
        "Docket Import ID": item.get("docket_import_id", ""),
        "Case Identity Checks": str(item.get("identity_match", {}).get("checks", {})),
        "Client Executive Summary": item.get("client_summary", ""),
        "Case Filing Date (Source)": item.get("date_filed", ""),
        "Termination Date": item.get("date_terminated", ""),
        "Detected Procedural Milestones": ", ".join(item.get("matched_keywords", [])),
        "Trigger Docket Event": item.get("trigger_entry", ""),
        "Trigger Date": item.get("trigger_date", ""),
        "Trigger Entry Number": item.get("trigger_entry_number", ""),
        "Trigger Date Basis": item.get("trigger_date_basis", ""),
        "Docket Source Link": item.get("trigger_source_url") or item.get("docket_url", ""),
        "PacerMonitor Docket Link": item.get("pacermonitor_url", ""),
        "PacerMonitor Search Link": item.get("search_url", ""),
        "Retrieved At (UTC)": item.get("retrieved_at", ""),
        "Source Last Updated": item.get("source_last_updated", ""),
        "Latest Retrieved Entry Date": item.get("latest_entry_date", ""),
        "Source Coverage": item.get("coverage", ""),
        "Undated Entries Excluded": item.get("undated_entry_count", 0),
        "Source Page SHA256": item.get("page_sha256", ""),
        "Source Errors": "; ".join(d["source"] + ": " + d["error"] for d in item.get("source_diagnostics", []) if d.get("error")),
    }

def generate_enriched_excel(results: List[Dict[str, Any]]) -> io.BytesIO:
    rows = []
    for item in results:
        row = dict(item.get("raw_row", {}))
        
        row.update(evidence_columns(item))

        rows.append(row)

    df = pd.DataFrame(rows)
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Patent Litigation Tracker")
        workbook = writer.book
        worksheet = writer.sheets["Patent Litigation Tracker"]
        
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Slate 900
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Emerald
        amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Amber
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")   # Rose
        teal_fill = PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")  # Teal for Invalidity
        purple_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid") # Purple for AI Verified
        gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            
        worksheet.row_dimensions[1].height = 30

        status_col_idx = None
        inv_col_idx = None
        verif_col_idx = None
        for idx, col_name in enumerate(df.columns, start=1):
            if col_name == "Status":
                status_col_idx = idx
            elif col_name == "Invalidity Contentions Observed?":
                inv_col_idx = idx
            elif col_name == "Verification Source":
                verif_col_idx = idx

        for row_num in range(2, len(df) + 2):
            worksheet.row_dimensions[row_num].height = 22
            status_val = str(worksheet.cell(row=row_num, column=status_col_idx).value or "") if status_col_idx else ""
            inv_val = str(worksheet.cell(row=row_num, column=inv_col_idx).value or "") if inv_col_idx else ""
            verif_val = str(worksheet.cell(row=row_num, column=verif_col_idx).value or "") if verif_col_idx else ""
            
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Style Status column
                if col_num == status_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if "Settled" in status_val:
                        cell.fill = amber_fill
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="92400E")
                    elif "Dismissed" in status_val:
                        cell.fill = red_fill
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
                    elif "Active" in status_val:
                        cell.fill = green_fill
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
                    else:
                        cell.fill = gray_fill

                # Style Invalidity Contentions column
                if col_num == inv_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if inv_val == "Yes":
                        cell.fill = teal_fill
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F766E")

                # Style Verification Source column
                if col_num == verif_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if "Gemini" in verif_val:
                        cell.fill = purple_fill
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="6B21A8")
                    else:
                        cell.fill = gray_fill

        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = min(len(val_str), 40)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output.seek(0)
    return output

def generate_csv(results: List[Dict[str, Any]]) -> str:
    rows = []
    for item in results:
        row = dict(item.get("raw_row", {}))
        row.update(evidence_columns(item))
        rows.append(row)
    df = pd.DataFrame(rows)
    return df.to_csv(index=False)
